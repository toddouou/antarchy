# -*- coding: utf-8 -*-
"""Generate the HIVE development working doc (.docx) from project knowledge."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x10, 0x12, 0x18)
ACCENT = RGBColor(0x5e, 0x60, 0xce)      # hive violet
MUTED = RGBColor(0x55, 0x5a, 0x66)
GOOD = RGBColor(0x1f, 0x7a, 0x4d)

doc = Document()

# ---- base styling -----------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK

for name, size, color in [("Heading 1", 18, ACCENT), ("Heading 2", 14, INK),
                          ("Heading 3", 11.5, ACCENT)]:
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = color
    st.font.bold = True


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hex_fill)
    tcPr.append(sh)


def set_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


def para(text="", *, bold=False, italic=False, color=None, size=None,
         align=None, space_after=6, style=None):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = color
        if size:
            r.font.size = Pt(size)
    return p


def runs(p, parts):
    """parts = list of (text, dict-of-attrs)."""
    for text, attrs in parts:
        r = p.add_run(text)
        r.bold = attrs.get("bold", False)
        r.italic = attrs.get("italic", False)
        if "color" in attrs:
            r.font.color.rgb = attrs["color"]
        if "size" in attrs:
            r.font.size = Pt(attrs["size"])
    return p


def bullet(text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)
    return p


def kv_table(rows, widths, header=None):
    n = len(rows[0])
    t = doc.add_table(rows=0, cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if header:
        hr = t.add_row().cells
        for i, h in enumerate(header):
            hr[i].text = ""
            rr = hr[i].paragraphs[0].add_run(h)
            rr.bold = True
            rr.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
            rr.font.size = Pt(9.5)
            shade(hr[i], "5e60ce")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            rr = cells[i].paragraphs[0].add_run(str(val))
            rr.font.size = Pt(9.5)
            if i == 0 and n > 2:
                rr.bold = True
    set_widths(t, widths)
    para(space_after=4)
    return t


def hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "c7c9d6")
    pbdr.append(bottom)
    pPr.append(pbdr)


def decision_table(rows):
    """rows = list of (feature, status, what-it-is). Adds blank KEEP? + NOTES cols."""
    t = doc.add_table(rows=0, cols=5)
    t.style = "Table Grid"
    header = ["Feature", "Now", "What it is", "Keep / Cut?", "Notes"]
    hr_cells = t.add_row().cells
    for i, h in enumerate(header):
        hr_cells[i].text = ""
        rr = hr_cells[i].paragraphs[0].add_run(h)
        rr.bold = True
        rr.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        rr.font.size = Pt(9)
        shade(hr_cells[i], "10121 8"[:6].replace(" ", "") or "101218")
        shade(hr_cells[i], "101218")
    for feat, status, what in rows:
        cells = t.add_row().cells
        vals = [feat, status, what, "", ""]
        for i, v in enumerate(vals):
            cells[i].text = ""
            rr = cells[i].paragraphs[0].add_run(v)
            rr.font.size = Pt(9)
            if i == 0:
                rr.bold = True
            if i == 1:
                rr.font.color.rgb = MUTED
    set_widths(t, [1.45, 0.6, 2.85, 0.95, 1.05])
    para(space_after=4)
    return t


# =============================================================================
# COVER
# =============================================================================
title = para("HIVE", bold=True, color=ACCENT, size=40, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Development Working Doc", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Planet-scale multiplayer Langton's-ant territory war", italic=True, color=MUTED,
     align=WD_ALIGN_PARAGRAPH.CENTER, size=11, space_after=14)
para("Living document — last assembled 2026-05-31", color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER,
     size=9, space_after=14)

hr()
para("How to use this doc", bold=True, size=12, space_after=4)
para("This is a shared brain for the two of us. Section 1 frames what HIVE is. Section 2 nails the "
     "core mechanics that are already real in the engine. Section 3 is everything else — shipped "
     "extras, half-built systems, and raw ideas — laid out so we can argue about what to keep.",
     space_after=6)
para("Every feature table has a Now column (is it in the code today?) plus blank Keep / Cut? and "
     "Notes columns. Fill them in together. When we decide something, log it in the Decisions section "
     "at the end so we stop re-litigating it.", space_after=8)

para("Status legend", bold=True, size=11, space_after=4)
kv_table([
    ["LIVE", "In the Rust engine right now, working."],
    ["WIP", "Partly built / wired but not finished or not validated in-browser."],
    ["IDEA", "On paper only — not started."],
    ["DEFER", "Deliberately parked; revisit later."],
    ["CUT", "Explicitly de-scoped — do not build (kept here for memory)."],
    ["DECIDE", "Needs a call from us before work starts."],
], widths=[0.9, 5.6], header=["Tag", "Meaning"])

doc.add_page_break()

# =============================================================================
# SECTION 1 — ABOUT THE GAME
# =============================================================================
doc.add_heading("1 · About the Game", level=1)

doc.add_heading("The pitch", level=2)
para("HIVE is a real-time, planet-scale multiplayer territory war built on Langton's ant. You drop a "
     "queen somewhere on a world map (real OpenStreetMap geography), then deploy worker ants that move "
     "by the classic Langton turn-and-flip rule. As they wander they paint territory in your color, and "
     "when they reach an enemy queen they besiege it. The whole thing plays out across a world the size "
     "of Earth — 1,500,000 × 750,000 tiles — that every player shares at once.", space_after=6)

doc.add_heading("Why it's interesting", level=2)
bullet("Emergent, not scripted. Langton's ant produces highways, chaos, and sudden order on its own — "
       "the math toy IS the gameplay surface. Placement and timing are the skill.")
bullet("Real geography. Your territory sits on your actual city. People can recognize home under the ant wars.")
bullet("One shared planet. Not matches — a persistent world where hundreds of queens fight at once.")
bullet("Screen-recordable. Glowing territory spreading across a map is inherently shareable (the viral thesis).")

doc.add_heading("Who it's for", level=2)
para("Casual browser-game players who like emergent strategy — the Slither.io / Generals.io / Diep.io "
     "crowd, but with a smarter cellular-automaton core. Easy to teach, deep to master, short clips to share.",
     space_after=6)

doc.add_heading("Where it stands today (H0)", level=2)
para("The live engine is HIVE-SIM, a Rust crate in sim/ (tokio + axum). It replaced an earlier Node.js "
     "prototype, which has been removed. What's running right now:", space_after=4)
bullet("Full Langton simulation on a dedicated OS thread at a configurable tick rate (default 50 Hz).")
bullet("Earth-sized world held in memory as a sparse grid of 256×256 tile chunks (only painted regions cost RAM).")
bullet("Canvas 2D client embedded in the binary: OpenStreetMap base map, Mercator projection, smooth "
       "pan/zoom, touch controls, fog-of-war, queen + ant rendering.")
bullet("WebSocket auth (SHA-256), live admin panel with sliders, leaderboard, shop/credits, prestige, regions, discovery.")
bullet("Viewport delivery decoupled onto its own thread and built in parallel (rayon) at ~25 Hz.")

para("Verified scale headroom: a 100k-ant tick benchmarks at ~10.4 ms — about 52% of the 20 ms budget "
     "at 50 Hz. The u16 palette-indexed tile store roughly halves territory RAM vs the old u32 layout.",
     italic=True, color=MUTED, space_after=6)

doc.add_heading("Known blockers before a public share", level=2)
bullet("No durable persistence — the world AND all accounts reset on every server restart.")
bullet("OSM tiles are fetched straight from tile.openstreetmap.org — against their policy at scale; "
       "needs a self-hosted / dedicated tile source before public load.")
bullet("Still Canvas 2D — no WebGL glow / shader territory layer yet.")
bullet("No viral capture/share hook yet (time-lapse, screenshot, spectate link).")

doc.add_heading("Product vision — the phases", level=2)
kv_table([
    ["H0", "DONE", "Rust engine, Earth-sized world, OSM map, full sim, admin, shop, regions."],
    ["H1", "IDEA", "Visual overhaul: WebGL2 renderer + MapLibre base map, neon/CRT look, viral capture."],
    ["H2", "IDEA", "Live server, real domain (HTTPS/WSS), TRUE persistence (WAL + snapshots), 100+ queens."],
    ["H3", "DESIGN", "Planet-scale rewrite: sharded multi-host sim, ScyllaDB tiles, 10k concurrent / 100k DAU."],
    ["H4", "IDEA", "Virality & retention: auto time-lapse bot, factions, seasonal resets, referrals, streamer widgets."],
], widths=[0.6, 0.8, 5.1], header=["Phase", "Status", "Target"])

doc.add_heading("Tech stack at a glance", level=2)
kv_table([
    ["Engine", "Rust — tokio + axum, rayon, serde, sha2, rustc-hash."],
    ["Sim model", "One process, one World behind Arc<RwLock>. Sim on a dedicated OS thread; WS tasks only push Cmds."],
    ["World", "1,500,000 × 750,000 tiles @ ~26.72 m/tile (≈ Earth's circumference). Sparse 256×256 chunks."],
    ["Tiles", "u16 palette-indexed chunks (Uniform / Dense), per-player counts kept in sync via tiles.set()."],
    ["Client", "Single-file public/client.html (Canvas 2D) — embedded via include_str!, so editing needs a recompile."],
    ["Wire", "Base64 LE-u16 tiles + base64-u8 fog, JSON messages over WebSocket. Rate-limited 120 msg/s/conn."],
    ["Persistence", "None yet — fresh empty world every boot; users.json is written but never read back."],
], widths=[1.1, 5.4], header=["Layer", "Choice"])

doc.add_page_break()

# =============================================================================
# SECTION 2 — CORE GAME MECHANICS
# =============================================================================
doc.add_heading("2 · Core Game Mechanics", level=1)
para("These are the load-bearing rules — what actually makes HIVE the game it is. All of this is LIVE in "
     "the engine. Where there's a number, it's the current default (and most are admin-tunable live; see "
     "the parameters table at the end of this section).", italic=True, color=MUTED, space_after=8)

doc.add_heading("The Langton's ant rule (the heart of everything)", level=2)
para("Each worker holds a position and a heading. Every tick it reads the tile it is standing on and acts:",
     space_after=4)
kv_table([
    ["Unclaimed (0)", "Turn counter-clockwise, then paint the tile your color, then step forward."],
    ["Your own color", "Turn clockwise, then step forward (your trails self-organize into highways)."],
    ["An enemy's color", "Go straight (no turn), step forward (you punch through enemy territory)."],
    ["World edge", "Reflect 180° and step back inward (no wrap)."],
], widths=[1.5, 5.0], header=["Tile under the ant", "What the ant does"])
para("This single rule produces the whole visual language: chaotic blooms near the queen, long straight "
     "highways once an ant locks into a pattern, and constant churn at the borders between colors.",
     space_after=6)

doc.add_heading("Queens", level=2)
bullet("Your one anchor. Placed on the map to enter the world; if it dies, your territory is wiped.")
bullet("Has HP (default base 100). Enemy ants reaching the footprint deal damage; your own ants touching it heal it (+1).")
bullet("Optional passive HP regen per second while below max (default 0 = off; admin-tunable).")
bullet("Footprint grows with level: 2×2 at low level up to 8×8 at level 100 (steps at 10/25/50/75/90/100).")
bullet("Queen positions are cached in a queen_map; flagged dirty whenever a queen is placed, killed, moved, or resized.")

doc.add_heading("Workers (ants)", level=2)
kv_table([
    ["Normal", "1×1, moves every tick, paints by the Langton rule, ×1 queen damage. The default worker."],
    ["Brute", "2×2, moves only on EVEN ticks (half speed), never erases its own tiles, ×3 queen damage. "
              "Bought from the shop. Intended as the heavy 'highway' / siege unit."],
], widths=[0.9, 5.6], header=["Type", "Behavior"])
bullet("Army cap: a player can field at most army_cap live workers at once (default 1000); deploy is blocked at the cap.")
bullet("Lifespan: workers die after lifespan ticks (default 4,320,000 ≈ 24h at 50 Hz). Forces ongoing production.")
bullet("Daily ant grant: non-NPC players get daily_ants (default 5) refilled each day; +levelup_ant_grant per level gained.")

doc.add_heading("Territory & tiles", level=2)
bullet("tiles.get(x,y) returns 0 (unclaimed) or the owning player's numeric id. There is no separate 'color' — owner IS color.")
bullet("Every mutation goes through tiles.set(x,y,owner), which keeps exact per-player counts in sync and drops empty chunks.")
bullet("Territory tile count is the backbone of score, leaderboard rank, and region holding.")

doc.add_heading("Clash resolution (5×5 majority convert)", level=2)
para("Where colors meet, the minority flips. For a contested cell, the engine tallies owners in the "
     "surrounding neighborhood; if one player holds at least convert_pct (default 0.65) of the area, ants "
     "of other colors on that cell convert to the dominant player and award convert XP. This is what makes "
     "fronts advance and lets a stronger force absorb a weaker one rather than just overlapping it.",
     space_after=6)

doc.add_heading("Combat — damage & healing", level=2)
bullet("Enemy ant on a queen footprint → queen takes ant_damage (default 1.0), brutes ×3. Damage is batched per queen per tick.")
bullet("Friendly ant on its own queen → heals +1 (capped at max HP).")
bullet("Damage tracks last attacker so kills can be attributed (killfeed / credits).")
bullet("Defender (shop): when an enemy worker comes within DEFENDER_RANGE (10 tiles) of your queen, a "
       "defender spawns just outside the footprint heading at the threat. Decays after 1h.")
bullet("Shield (shop): 12h queen invulnerability window.")

doc.add_heading("Fog of war", level=2)
para("The server computes a per-player fog distance field (two-pass Chebyshev distance transform): clear "
     "out to ~20 tiles, gradient out to ~30, then dark. Admins receive an all-zero field (no fog). Fog is "
     "computed on a padded viewport slice off the lock so it doesn't stall the tick.", space_after=6)

doc.add_heading("XP, levels & queen growth", level=2)
para("XP is queued during a tick (award_xp) and applied at flush. Sources and the curve:", space_after=4)
kv_table([
    ["Kill a queen", "xp_kill = 5000"],
    ["Convert an enemy ant", "xp_convert = 5"],
    ["Tile milestone", "xp_tile_award = 25 every xp_tile_milestone = 500 tiles painted"],
    ["Heal touch", "xp_heal = 1"],
    ["Hit a queen", "+0.5 per hit"],
    ["Curve", "total_xp_for_level(n) = xp_base·(n−1)^xp_exp  →  500·(n−1)^2.2 ; cap level 100"],
], widths=[1.9, 4.6], header=["Source", "Default"])
para("Level drives queen footprint size (see Queens). Note: the roadmap flags dropping the 'LEVEL X / 100' "
     "cap display in the HUD — the level bar already shows progress, not '/100'.", italic=True, color=MUTED, space_after=6)

doc.add_heading("Score formula", level=2)
para("calc_score = tiles + 0.5 · seconds_alive + 500 · kills. This (Grand Score), not raw tile count, is "
     "what the leaderboard now sorts by.", space_after=6)

doc.add_heading("Credits, shop & death/prestige", level=2)
bullet("Credits are earned ONLY by killing an enemy queen (+1 each), hard-capped at 100 (CREDIT_CAP).")
bullet("Shop spends credits on the items in §3. Highway / relocate / brute are placement items; "
       "defender / shield / alliance are effects.")
bullet("On queen death: territory disappears and most state resets — EXCEPT prestige, lifetime stats, "
       "credits, and discovery (countries visited), which persist. Redeploying increases prestige.")
bullet("NPC queens (p.npc = true) are excluded from daily refills and normal leaderboard behavior.")
bullet("Season: when uptime passes season_secs (default 30 days), the world auto-wipes and a fresh season starts.")

doc.add_heading("Tunable parameters (admin sliders, live)", level=2)
para("All of these live in the global Config and are mutated live via the admin panel, clamped to the "
     "ranges shown. This is the dials board — most balance experiments happen here without a recompile.",
     space_after=4)
kv_table([
    ["tick_rate", "50", "1 – 500", "Sim Hz."],
    ["lifespan", "4,320,000", "1k – 8.64M", "Ticks a worker lives (~24h @ 50 Hz)."],
    ["bubble_r", "30", "5 – 5,000", "Deploy radius around the queen."],
    ["hp_base", "100", "1 – 1,000,000", "Queen base HP."],
    ["convert_pct", "0.65", "0.1 – 1.0", "Majority fraction needed to convert a clash."],
    ["daily_ants", "5", "0 – 1,000", "Daily worker refill."],
    ["ant_damage", "1.0", "0.1 – 50", "Damage per enemy-ant hit (brute ×3)."],
    ["hp_regen", "0.0", "0 – 100", "Passive queen HP/sec below max."],
    ["army_cap", "1000", "1 – 1,000,000", "Max live workers per player."],
    ["xp_base", "500", "1 – 1,000,000", "Level-curve base."],
    ["xp_exp", "2.2", "0.5 – 5", "Level-curve exponent."],
    ["xp_kill", "5000", "0 – 1,000,000", "XP per queen kill."],
    ["xp_convert", "5", "0 – 10,000", "XP per ant converted."],
    ["xp_tile_milestone", "500", "1 – 100,000", "Tiles per milestone award."],
    ["xp_tile_award", "25", "0 – 10,000", "XP per milestone."],
    ["levelup_ant_grant", "1", "0 – 100", "Bonus workers per level gained."],
    ["spawn_pan", "200", "10 – 10,000", "Initial camera pan span."],
    ["season_secs", "2,592,000", "0 – 31.5M", "Season length; 0 = no auto-wipe."],
], widths=[1.55, 1.0, 1.3, 2.65], header=["Param", "Default", "Range", "What it controls"])

doc.add_page_break()

# =============================================================================
# SECTION 3 — OTHER MECHANICS & FEATURES
# =============================================================================
doc.add_heading("3 · Other Mechanics & Features", level=1)
para("Everything around the core loop: shipped extras, half-built systems, the roadmap, and parked ideas. "
     "Use the Keep / Cut? and Notes columns to make calls together. 'Now' uses the legend on page 1.",
     italic=True, color=MUTED, space_after=8)

doc.add_heading("3.1 · Shop items", level=2)
para("The shop is LIVE (credits-only economy, §2). Current catalog:", space_after=4)
decision_table([
    ["Highway", "LIVE", "Buy a diagonal painted road (HIGHWAY_LEN 150 tiles); start must be within 30 tiles of your territory. PRICE 10."],
    ["Relocate", "LIVE", "Move your queen to a new spot. PRICE 20."],
    ["Defender", "LIVE", "Auto-spawns a guard ant when an enemy worker enters 10 tiles; decays after 1h. PRICE 1."],
    ["Brute", "LIVE", "Deploy a 2×2 half-speed ×3-damage siege worker. PRICE 20."],
    ["Shield", "LIVE", "12h queen invulnerability. PRICE 10."],
    ["Alliance", "WIP", "Slot + price (80) exist but shop-buy is a no-charge 'coming soon' stub. See 3.6."],
])
para("Roadmap note: 'fix highways' — the highway purchase and the brute 'highway' behavior both want a "
     "tuning/repair pass (white-surface mechanics + 2×2 movement), not a rebuild. Brute is already 2×2 / "
     "even-tick / ×3.", italic=True, color=MUTED, space_after=6)

doc.add_heading("3.2 · Interface / UX overhaul", level=2)
para("UX is the stated top priority. This slice was specced and largely built in compile-verified "
     "checkpoints (CP1–CP4); most is wired and server-verified, with in-browser visual validation still "
     "pending in places.", space_after=4)
decision_table([
    ["Header = server stats", "LIVE", "Header shows online players, live ants, live queens, tick/TPS/uptime (folds in server-health). Personal score removed from header."],
    ["Logout top-right", "LIVE", "Moved out of the sidebar into the header's top-right."],
    ["Killfeed (CS1.6 style)", "LIVE", "Top-right vertical stack, newest on top, fades ~6s, color-coded: '♛ Killer ⚔ Victim'. Driven by a structured server kill event."],
    ["Queen + Stats merged", "WIP", "Combined into one sidebar panel with a STATS sub-divider. Restyle (spacing/type) still pending a look."],
    ["Workers panel redo", "LIVE", "Type picker (normal/brute), READY/ARMY/cap/REFILL grid + army bar, expandable active-ant list with lifespan-countdown bars."],
    ["Leaderboard overhaul", "LIVE", "Sorts by Grand Score, region column + per-region tabs, pins MY row when outside top-N, EXPAND → top-30 + LV/K/score/KD."],
    ["Target/inspect window", "WIP", "Folded into the compass coordinate overlay; admin TARGET pane untouched. Appearance polish outstanding."],
    ["Login patch-notes", "IDEA", "A 'what's new' panel on login. Not built."],
])

doc.add_heading("3.3 · Regions — the 'pride' / metro system", level=2)
para("Data-driven, fightable metros plus a country fallback, so every spot on Earth has an identity to "
     "fight over. LIVE and wire-verified.", space_after=4)
bullet("10 starter metros ship in data/regions.json: New York, Los Angeles, Paris, London, Tokyo, "
       "São Paulo, Mexico City, Lagos, Mumbai, Berlin. Each is lat/lon + radius_km.")
bullet("Inside a metro radius → that metro. Outside all metros → country tag (point-in-polygon vs Natural "
       "Earth countries geojson, embedded in the binary). Neither → 'Open Water'.")
bullet("HOLDER = whoever has the most painted tiles inside the radius (king-of-the-hill by area, recomputed "
       "live every ~500 ticks, strided for sparse-safety). Shown as 'HELD BY <name>' in the region switcher + leaderboard.")
bullet("Region shown in the header switcher (click a metro → camera flies there) and as a leaderboard column/tab.")
bullet("DE-SCOPED by user: drawing region rings/borders on the map.")
para("Idea space: more metros ship as updates; continents are already derivable (used by discovery).",
     italic=True, color=MUTED, space_after=6)

doc.add_heading("3.4 · Discovery (countries & continents)", level=2)
decision_table([
    ["Passport tracking", "LIVE", "Account-level visited_countries / visited_continents sets that SURVIVE queen death. Sampled from roaming ants + queen placement via the geojson."],
    ["Discovery window", "LIVE", "Sidebar ◆ DISCOVERY button → window showing the passport + a USER-vs-QUEEN stat split."],
    ["User-vs-queen stats", "LIVE", "Lifetime kills, lifetime peak tiles, queens fielded carried on 'me'."],
])

doc.add_heading("3.5 · Welcome-back & death screen", level=2)
decision_table([
    ["Welcome-back summary", "LIVE", "On reconnect after >30s away (with a live queen), diffs an away-snapshot → rewarding feed: ants converted/lost, kills, tiles gained/lost, levels. Queen-died path leads with 'queen fell'."],
    ["Death screen", "LIVE", "On queen death: summarizes the queen's life (peak tiles, kills, time alive, level, region, credits), explains prestige, REDEPLOY button (redeploy bumps prestige). Ships WITHOUT countries-visited line for now."],
])

doc.add_heading("3.6 · Parked / de-scoped systems", level=2)
para("Explicitly set aside by the user. Kept here so we remember the decision and the design.", space_after=4)
decision_table([
    ["Alliance system", "DEFER", "Buy a key in shop → create / request / invite. Needs alliance state on Player, no-friendly-fire + shared-vision rules, and UI. Currently a no-charge stub. De-scoped for now."],
    ["Ant tracking / cinema cam", "DEFER", "A follow-cam ('cinema_coefficient') that tracks a chosen ant or queen with smoothed pan/zoom. De-scoped for now."],
    ["Region rings / borders", "CUT", "Drawing metro/region boundaries on the map. Explicitly cut — holders show as text only."],
    ["Playtest feedback (Djani / Emily / Amina / France)", "DEFER", "Notes exist but details not captured. Ask the user before implementing."],
])

doc.add_heading("3.7 · Scaling & planet-scale (the long game)", level=2)
para("The optimization effort is far along — P1–P5 are done. Targets: ~100k ants, 500 queens, activity in "
     "all 4 corners, 1-month seasons, ~100 billion painted cells at peak (oceans never painted).",
     space_after=4)
decision_table([
    ["Decoupled viewport thread", "LIVE", "Viewport serialization moved off the sim thread → steady tick cadence (the smoothness fix)."],
    ["u16 palette tile store", "LIVE", "Uniform/Dense chunks, recycled palette indices → roughly half the territory RAM; the planet-scale memory fix."],
    ["Parallel sim", "LIVE", "Move-plan + periodic ant sorts on rayon; 100k-ant tick ~10.4 ms (52% of budget). Deep chunk-partitioned paint deliberately NOT added (unsafe + unnecessary)."],
    ["LOD territory pyramid", "WIP", "Server downsamples to a ≤800-cell grid when zoomed out (wire-verified); actual overview pixels need browser validation."],
    ["Season ops + /health metrics", "LIVE", "Periodic compaction, tile/chunk/tick metrics on /health, auto season wipe."],
    ["u16 indices on the wire (P2b)", "IDEA", "Send palette indices instead of clamped ids; fixes the >65535 clamp bug. Needs client palette-by-index changes."],
    ["Sharded multi-host (H3)", "IDEA", "1024×1024 region shards, cross-shard ant migration, ScyllaDB tiles, Go coordinator. Design-only until H2 traction."],
])

doc.add_heading("3.8 · Visual overhaul & viral hooks (H1 / H4)", level=2)
decision_table([
    ["WebGL2 renderer", "IDEA", "Replace Canvas 2D: territory as a fullscreen shader, additive-blend ant glow, queen pulse shader. ~100× faster, smooth zoom."],
    ["MapLibre GL base map", "IDEA", "Real vector/raster base map behind the grid; Google-Maps-feel panning; recognizable cities."],
    ["Neon + CRT aesthetic", "IDEA", "Preserve the MYRMIDON look: scanlines, neon palette, pixel-aligned glow."],
    ["Time-lapse capture", "IDEA", "Client records last ~300 ticks → export MP4/GIF. The core viral hook."],
    ["Territory screenshot", "IDEA", "One-click shareable PNG with stats overlay (territory %, level, name)."],
    ["Spectate link", "IDEA", "?spectate=playerid opens the game panned to that queen."],
    ["Auto time-lapse bot (H4)", "IDEA", "Server auto-posts a daily highlight reel of the biggest war."],
    ["Factions / referral / streamer widget (H4)", "IDEA", "Retention layer — post-traction."],
])

doc.add_heading("3.9 · Persistence (H2 blocker)", level=2)
decision_table([
    ["World survives restart", "IDEA", "WAL of changed tiles per tick + full snapshot every ~15 min; replay journal over last snapshot on boot."],
    ["Account persistence", "IDEA", "users.json is written today but never read back; accounts reset each boot. Wire load on startup."],
    ["Real domain HTTPS/WSS", "IDEA", "Deploy target ~Hetzner CPX31; scales to ~500 concurrent before sharding (H3)."],
])

doc.add_heading("3.10 · Monetization & game-mode ideas (MYRMIDON archive)", level=2)
para("From the earlier 1v1-arena design (MYRMIDON), before the open-world pivot. Re-evaluate which of these "
     "still make sense for a persistent shared planet — most need rethinking for HIVE, but the thinking is here.",
     italic=True, color=MUTED, space_after=4)
decision_table([
    ["Ads (between-match / sidebar)", "IDEA", "Designed around discrete matches — doesn't map cleanly to a persistent world. Rethink for HIVE."],
    ["Premium tiers (ARCHON / lifetime)", "IDEA", "No-ads, cosmetics, custom rules, priority. Free-to-play, never pay-to-win was the rule."],
    ["Custom Langton rule strings", "IDEA", "RLR / LRRRRRLLR etc. as premium unlocks. Interesting depth lever even in HIVE."],
    ["Cosmetic skins (queen/ant/trail)", "IDEA", "Clean monetization that fits HIVE."],
    ["Ant variants (Soldier/Scout/Drone)", "IDEA", "We already shipped Normal + Brute; this is the same axis to extend."],
    ["Power-up cells on the map", "IDEA", "Pickups: extra spawn, shield, paint burst. Could scatter across the planet."],
    ["Alt win/objective modes (KotH, capture-the-queen, fog)", "IDEA", "Arena-era modes; the region holder system is already a KotH-by-area in spirit."],
    ["Achievements / replays / tournaments", "IDEA", "Deep retention; far out for HIVE."],
])

doc.add_page_break()

# =============================================================================
# DECISIONS + OPEN QUESTIONS + PARKING LOT
# =============================================================================
doc.add_heading("4 · Decisions Log", level=1)
para("When we settle something, write it here so we stop re-arguing it. Date it.", italic=True, color=MUTED, space_after=6)
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
for i, h in enumerate(["Date", "Decision", "Why"]):
    t.rows[0].cells[i].text = ""
    rr = t.rows[0].cells[i].paragraphs[0].add_run(h)
    rr.bold = True
    rr.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
    shade(t.rows[0].cells[i], "101218")
for _ in range(8):
    t.add_row()
set_widths(t, [0.9, 3.3, 2.3])
para(space_after=8)

doc.add_heading("5 · Open Questions", level=1)
bullet("What's the #1 thing to build next — UX polish slice, persistence, or the WebGL/viral visual pass?")
bullet("Do we ship persistence before any public share, or run disposable seasons on purpose?")
bullet("Self-hosted map tiles: when, and which provider?")
bullet("Does monetization belong in HIVE at all, or is this an open/free passion build for now?")
bullet("Brute/highway: what specifically feels wrong today — speed, damage, paint behavior, or cost?")
bullet("Alliance: revive it, or keep it parked? If revived, friendly-fire + shared-vision rules need design.")
bullet("Which playtest feedback (Djani / Emily / Amina / France) is still worth chasing?")

doc.add_heading("6 · Parking Lot — raw ideas", level=1)
para("Dump anything here, no judgment. Promote the good ones into §3 with a status tag.", italic=True, color=MUTED, space_after=6)
for _ in range(10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.add_run("•  ").font.color.rgb = ACCENT

out = r"C:\Users\Todd\OneDrive\Desktop\langton's ant\HIVE - Development Doc.docx"
doc.save(out)
print("SAVED:", out)
